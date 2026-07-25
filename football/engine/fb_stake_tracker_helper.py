"""
football/engine/fb_stake_tracker_helper.py
------------------------------------------
Helper functions for manual stake logging via checkboxes in opportunity lists.
"""

import os
import sys
import json
import csv
import hashlib
import re
import threading
from datetime import datetime

# Reconfigure stdout to UTF-8 to prevent encoding errors on Windows
if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass
import openpyxl

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
STAKED_HISTORY_FILE = os.path.join(_ROOT, 'data', 'staked_history.json')
ACTIVE_OPPS_FILE = os.path.join(_ROOT, 'data', 'active_opportunities.json')
BAL_UNBAL_STAKE_TRACKER_FILE = os.path.join(_ROOT, 'data', 'balanced_unbalanced_stake_tracker.xlsx')
LEGACY_STAKE_TRACKER_FILE = os.path.join(_ROOT, 'data', 'stake_tracker.xlsx')
QUASI_STAKE_TRACKER_FILE = os.path.join(_ROOT, 'data', 'quasi_stake_tracker.xlsx')

_file_lock = threading.RLock()

HEADERS = [
    'Opportunity ID', 'Category', 'Date ', 'Time', 'Match', 'Tournament', 'Kickoff', 'Market',
    'Platform 1', 'Bet 1', 'Odds 1', 'Stake 1 (GHS)',
    'Platform 2', 'Bet 2', 'Odds 2', 'Stake 2 (GHS)',
    'Platform 3', 'Bet 3', 'Odds 3', 'Stake 3 (GHS)',
    'Total Stake (GHS)', ' Profit %', ' Profit(GHS)', 'Accumulated Profit (GHS)'
]

QUASI_HEADERS = [
    'Opportunity ID', 'Category', 'Date', 'Time', 'Match', 'Tournament', 'Kickoff', 'Market',
    'Break-Even Outcome',
    'Best Profit Outcome',
    'Platform 1', 'Bet 1', 'Odds 1', 'Stake 1 (GHS)', 'Potential Profit If Wins 1 (GHS)',
    'Platform 2', 'Bet 2', 'Odds 2', 'Stake 2 (GHS)', 'Potential Profit If Wins 2 (GHS)',
    'Platform 3', 'Bet 3', 'Odds 3', 'Stake 3 (GHS)', 'Potential Profit If Wins 3 (GHS)',
    'Total Stake (GHS)',
    'Best Profit %',
    'Best Profit (GHS)',
    'Actual Winning Outcome',
    'Profit Won (GHS)',
    'Stake Returned? (yes/no)',
    'Result Notes',
]


def get_opp_id(opp):
    """Generates a stable 8-character hash for an opportunity."""
    match = opp.get('match', '')
    kickoff = opp.get('kickoff', '')
    market = opp.get('market', '')
    category = opp.get('category', '')
    
    # Sort platform details so ID is stable across ordering
    bets = sorted(opp.get('bets', []), key=lambda x: x.get('platform', ''))
    bets_str = "".join(f"{b.get('platform','')}_{b.get('outcome','')}_{b.get('odds',0)}_{b.get('stake',0)}" for b in bets)
    
    raw = f"{match}|{kickoff}|{market}|{category}|{bets_str}"
    return hashlib.md5(raw.encode('utf-8')).hexdigest()[:8]


def load_staked_history():
    """Loads staked history from JSON file."""
    if not os.path.exists(STAKED_HISTORY_FILE):
        return set()
    try:
        with open(STAKED_HISTORY_FILE, 'r', encoding='utf-8') as f:
            data = json.load(f)
            return set(data)
    except Exception as e:
        print(f"  WARNING [Stake Logger] Error loading staked history: {e}")
        return set()


def save_staked_history(history):
    """Saves staked history to JSON file."""
    try:
        os.makedirs(os.path.dirname(STAKED_HISTORY_FILE), exist_ok=True)
        with open(STAKED_HISTORY_FILE, 'w', encoding='utf-8') as f:
            json.dump(list(history), f, indent=2)
    except Exception as e:
        print(f"  WARNING [Stake Logger] Error saving staked history: {e}")


def load_active_opportunities():
    """Loads active opportunities cache."""
    if not os.path.exists(ACTIVE_OPPS_FILE):
        return {}
    try:
        with open(ACTIVE_OPPS_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    except Exception as e:
        print(f"  WARNING [Stake Logger] Error loading active opportunities: {e}")
        return {}


def save_active_opportunities(opps_dict):
    """Saves active opportunities cache."""
    try:
        os.makedirs(os.path.dirname(ACTIVE_OPPS_FILE), exist_ok=True)
        with open(ACTIVE_OPPS_FILE, 'w', encoding='utf-8') as f:
            json.dump(opps_dict, f, indent=2)
    except Exception as e:
        print(f"  WARNING [Stake Logger] Error saving active opportunities: {e}")


def ensure_stake_tracker():
    """Ensures the balanced/unbalanced stake tracker exists, migrating old names if needed."""
    if os.path.exists(BAL_UNBAL_STAKE_TRACKER_FILE):
        _ensure_workbook_headers(BAL_UNBAL_STAKE_TRACKER_FILE, "Stakes", HEADERS)
        return
        
    os.makedirs(os.path.dirname(BAL_UNBAL_STAKE_TRACKER_FILE), exist_ok=True)

    if os.path.exists(LEGACY_STAKE_TRACKER_FILE):
        os.replace(LEGACY_STAKE_TRACKER_FILE, BAL_UNBAL_STAKE_TRACKER_FILE)
        return
    
    # Check if we should migrate from existing CSV
    old_csv_path = LEGACY_STAKE_TRACKER_FILE.replace('.xlsx', '.csv')
    if os.path.exists(old_csv_path):
        print("  Migration [Migration] Migrating data/stake_tracker.csv to data/balanced_unbalanced_stake_tracker.xlsx...")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Stakes"
        
        try:
            with open(old_csv_path, 'r', encoding='utf-8') as f:
                reader = csv.reader(f)
                for row in reader:
                    ws.append(row)
            wb.save(BAL_UNBAL_STAKE_TRACKER_FILE)
            print("  OK [Migration] Migration completed successfully.")
            # Rename old CSV file to keep a backup
            try:
                os.rename(old_csv_path, old_csv_path + ".bak")
            except Exception:
                pass
            return
        except Exception as e:
            print(f"  WARNING [Migration] Error migrating CSV to XLSX: {e}")
            
    # Default: create fresh empty workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stakes"
    ws.append(HEADERS)
    wb.save(BAL_UNBAL_STAKE_TRACKER_FILE)


def _ensure_workbook_headers(path, sheet_title, headers):
    """Create a workbook if missing, and append any missing headers to row 1."""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    if not os.path.exists(path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_title
        ws.append(headers)
        wb.save(path)
        return

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    existing = [cell.value for cell in ws[1]]
    changed = False
    for header in headers:
        if header not in existing:
            ws.cell(row=1, column=len(existing) + 1, value=header)
            existing.append(header)
            changed = True
    if changed:
        wb.save(path)


def ensure_quasi_stake_tracker():
    """Ensures the quasi stake tracker exists for manually settled quasi bets."""
    _ensure_workbook_headers(QUASI_STAKE_TRACKER_FILE, "Quasi Stakes", QUASI_HEADERS)


def _format_kickoff(kickoff_str):
    try:
        for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M'):
            try:
                dt_ko = datetime.strptime(str(kickoff_str).strip(), fmt)
                return f"{dt_ko.month}/{dt_ko.day}/{dt_ko.year} {dt_ko.strftime('%H:%M')}"
            except ValueError:
                continue
    except Exception:
        pass
    return kickoff_str or ''


def log_to_stake_tracker(opp, date_now):
    """Logs a balanced/unbalanced staked opportunity preserving Excel structures."""
    ensure_stake_tracker()
    
    bets = opp.get('bets', [])
    p1 = bets[0] if len(bets) > 0 else {}
    p2 = bets[1] if len(bets) > 1 else {}
    p3 = bets[2] if len(bets) > 2 else {}
    
    # Format date: M/D/YYYY
    date_str = f"{date_now.month}/{date_now.day}/{date_now.year}"
    time_str = date_now.strftime("%H:%M:%S")
    
    kickoff_str = _format_kickoff(opp.get('kickoff', ''))
        
    total_staked = sum(float(b.get('stake', 0)) for b in bets)
    
    profit_pct = opp.get('profit_pct', 0.0)
    category = opp.get('category', 'balanced')
    
    if category == 'quasi':
        profit_ghs = opp.get('best_profit_ghs', 0.0)
    elif category == 'unbalanced':
        profit_ghs = opp.get('max_profit_ghs', 0.0)
    else:
        profit_ghs = opp.get('profit_ghs', 0.0)
        
    row_values = [
        get_opp_id(opp),
        category,
        date_str,
        time_str,
        opp.get('match', ''),
        opp.get('tournament', ''),
        kickoff_str,
        opp.get('market', ''),
        p1.get('platform', ''),
        p1.get('outcome', ''),
        p1.get('odds', ''),
        round(float(p1.get('stake', 0)), 2) if p1.get('stake') else '',
        p2.get('platform', ''),
        p2.get('outcome', ''),
        p2.get('odds', ''),
        round(float(p2.get('stake', 0)), 2) if p2.get('stake') else '',
        p3.get('platform', ''),
        p3.get('outcome', ''),
        p3.get('odds', '') if p3 else '',
        round(float(p3.get('stake', 0)), 2) if p3.get('stake') else '',
        round(total_staked, 2),
        round(float(profit_pct), 2),
        round(float(profit_ghs), 2),
        ''  # Accumulated Profit is left blank for Excel Table formula auto-fill
    ]
    
    wb = openpyxl.load_workbook(BAL_UNBAL_STAKE_TRACKER_FILE)
    ws = wb.active
    ws.append(row_values)
    
    # Expand Table reference ranges if Excel Tables are active
    if ws.tables:
        for table in list(ws.tables.values()):
            ref = table.ref
            start_cell, end_cell = ref.split(':')
            match = re.match(r"([A-Z]+)([0-9]+)", end_cell)
            if match:
                col = match.group(1)
                new_ref = f"{start_cell}:{col}{ws.max_row}"
                table.ref = new_ref
                
    wb.save(BAL_UNBAL_STAKE_TRACKER_FILE)


def _expand_excel_tables(ws):
    """Expand any Excel table references after appending rows."""
    if not ws.tables:
        return
    for table in list(ws.tables.values()):
        ref = table.ref
        start_cell, end_cell = ref.split(':')
        match = re.match(r"([A-Z]+)([0-9]+)", end_cell)
        if match:
            col = match.group(1)
            table.ref = f"{start_cell}:{col}{ws.max_row}"


def _best_profit_outcome(opp):
    if opp.get('best_outcome'):
        return opp.get('best_outcome', '')
    bets = opp.get('bets', [])
    if not bets:
        return ''
    best_bet = max(
        bets,
        key=lambda b: float(b.get('profit_if_wins', 0) or 0)
    )
    return best_bet.get('outcome', opp.get('best_outcome', ''))


def _break_even_outcome(opp):
    if opp.get('break_even_outcome'):
        return opp.get('break_even_outcome', '')
    bets = opp.get('bets', [])
    if not bets:
        return ''
    break_even = min(
        bets,
        key=lambda b: abs(float(b.get('profit_if_wins', 0) or 0))
    )
    return break_even.get('outcome', '')


def log_to_quasi_stake_tracker(opp, date_now):
    """Logs a staked quasi opportunity to quasi_stake_tracker.xlsx."""
    ensure_quasi_stake_tracker()

    bets = opp.get('bets', [])
    p1 = bets[0] if len(bets) > 0 else {}
    p2 = bets[1] if len(bets) > 1 else {}
    p3 = bets[2] if len(bets) > 2 else {}

    date_str = f"{date_now.month}/{date_now.day}/{date_now.year}"
    time_str = date_now.strftime("%H:%M:%S")
    kickoff_str = _format_kickoff(opp.get('kickoff', ''))
    total_staked = sum(float(b.get('stake', 0)) for b in bets)
    best_profit = float(opp.get('best_profit_ghs', opp.get('profit_ghs', 0)) or 0)
    best_profit_pct = float(opp.get('profit_pct', 0) or 0)
    category = 'quasi'

    row_values = [
        get_opp_id(opp),
        category,
        date_str,
        time_str,
        opp.get('match', ''),
        opp.get('tournament', ''),
        kickoff_str,
        opp.get('market', ''),
        _break_even_outcome(opp),
        _best_profit_outcome(opp),
        p1.get('platform', ''),
        p1.get('outcome', ''),
        p1.get('odds', ''),
        round(float(p1.get('stake', 0)), 2) if p1.get('stake') else '',
        round(float(p1.get('profit_if_wins', 0)), 2) if p1.get('profit_if_wins') not in (None, '') else '',
        p2.get('platform', ''),
        p2.get('outcome', ''),
        p2.get('odds', ''),
        round(float(p2.get('stake', 0)), 2) if p2.get('stake') else '',
        round(float(p2.get('profit_if_wins', 0)), 2) if p2.get('profit_if_wins') not in (None, '') else '',
        p3.get('platform', ''),
        p3.get('outcome', ''),
        p3.get('odds', '') if p3 else '',
        round(float(p3.get('stake', 0)), 2) if p3.get('stake') else '',
        round(float(p3.get('profit_if_wins', 0)), 2) if p3.get('profit_if_wins') not in (None, '') else '',
        round(total_staked, 2),
        round(best_profit_pct, 2),
        round(best_profit, 2),
        '',
        '',
        '',
        '',
    ]

    wb = openpyxl.load_workbook(QUASI_STAKE_TRACKER_FILE)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    row = {header: value for header, value in zip(QUASI_HEADERS, row_values)}
    ws.append([row.get(header, '') for header in headers])
    _expand_excel_tables(ws)
    wb.save(QUASI_STAKE_TRACKER_FILE)


def check_and_log_ticked_bets():
    """Checks opportunity files for [s]/[S] stake markers, logs them, and updates files."""
    with _file_lock:
        data_dir = os.path.join(_ROOT, 'data')
        files = {
            'balanced': os.path.join(data_dir, 'intensive_balanced.txt'),
            'unbalanced': os.path.join(data_dir, 'intensive_unbalanced.txt'),
            'quasi': os.path.join(data_dir, 'intensive_quasi.txt')
        }
        
        staked_history = load_staked_history()
        active_opps = load_active_opportunities()
        
        logged_any = False
        logged_files = set()
        
        for cat, filepath in files.items():
            if not os.path.exists(filepath):
                continue
                
            try:
                with open(filepath, 'r', encoding='utf-8') as f:
                    content = f.read()
            except Exception as e:
                print(f"  WARNING [Stake Logger] Error reading {filepath}: {e}")
                continue
                
            # Matches [s] or [S] stake marker and extracts the 8-char hex ID
            pattern = r"\[[sS]\]\s+STAKE\s+THIS\s+OPP\s+\(ID:\s*([a-f0-9]+)\)"
            checked_ids = re.findall(pattern, content)
            
            if not checked_ids:
                continue
                
            modified_content = content
            
            for opp_id in checked_ids:
                if opp_id in staked_history:
                    # Update label in file if it wasn't modified to visual status yet
                    modified_content = re.sub(
                        rf"(?:ALERT\s+)?\[[sS]\]\s+STAKE\s+THIS\s+OPP\s+\(ID:\s*{opp_id}\)",
                        f"OK [s] STAKED (ID: {opp_id})",
                        modified_content
                    )
                    continue
                    
                opp = active_opps.get(opp_id)
                if opp:
                    now = datetime.now()
                    try:
                        category = opp.get('category', cat)
                        if category == 'quasi' or cat == 'quasi':
                            log_to_quasi_stake_tracker(opp, now)
                            logged_files.add("football/data/quasi_stake_tracker.xlsx")
                        else:
                            log_to_stake_tracker(opp, now)
                            logged_files.add("football/data/balanced_unbalanced_stake_tracker.xlsx")
                        staked_history.add(opp_id)
                        logged_any = True
                        print(f"Done [Stake Logger] Logged staked bet: {opp.get('match')} | {opp.get('market')} | ID: {opp_id}")
                        modified_content = re.sub(
                            rf"(?:ALERT\s+)?\[[sS]\]\s+STAKE\s+THIS\s+OPP\s+\(ID:\s*{opp_id}\)",
                            f"OK [s] STAKED (ID: {opp_id})",
                            modified_content
                        )
                    except PermissionError:
                        print(f"WARNING [Stake Logger] ERROR: Cannot write to the stake tracker because it is locked (likely open in Excel). Close Excel so Python can save it. Bet {opp_id} remains unchecked as STAKED in state.")
                    except Exception as ex:
                        print(f"ERROR [Stake Logger] Failed to log bet {opp_id}: {ex}")
                else:
                    print(f"WARNING [Stake Logger] Warning: Opportunity ID {opp_id} not found in active opportunities.")
                    
            if modified_content != content:
                try:
                    with open(filepath, 'w', encoding='utf-8') as f:
                        f.write(modified_content)
                except Exception as e:
                    print(f"  WARNING [Stake Logger] Error writing back to {filepath}: {e}")
                    
        if logged_any:
            save_staked_history(staked_history)
            
    # Push to GitHub outside the lock so network operations do not block local file access
    if logged_any:
        try:
            from engine.fb_arb_tracker import push_to_github
            push_to_github(
                filepaths=sorted(logged_files),
                message="Manual stake logged via text checkbox",
                quiet=True
            )
        except Exception as git_err:
            print(f"  WARNING [Git Sync] Error syncing stake tracker files: {git_err}")
